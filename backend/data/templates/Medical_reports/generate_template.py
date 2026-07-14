from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create a new document
doc = Document()

# 设置默认字体
style = doc.styles['Normal']
style.font.name = '仿宋_GB2312'
style.font.size = Pt(12)

# 标题 - 居中加粗
title = doc.add_paragraph()
title_run = title.add_run('{{ hospital_name }}X线摄影报告单')
title_run.font.size = Pt(16)
title_run.font.bold = True
title_run.font.name = '黑体'
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 患者基本信息 - 一行显示
info_para = doc.add_paragraph()
info_run = info_para.add_run(
    '姓名：{{ patient_info.name }}       '
    '性别：{{ patient_info.gender }} '
    '年龄：{{ patient_info.age }} '
    'X片号：{{ patient_info.xray_id }}'
)
info_run.font.size = Pt(12)
info_run.font.name = '仿宋_GB2312'

# 检查日期和部位
exam_para = doc.add_paragraph()
exam_run = exam_para.add_run(
    '检查日期：{{ examination_date }}      '
    '检查部位：{{ examination_site }}                        '
    '图象所见：'
)
exam_run.font.size = Pt(12)
exam_run.font.name = '仿宋_GB2312'

# 空行
doc.add_paragraph()

# 图象所见内容 - 左对齐
findings_para = doc.add_paragraph()
findings_run = findings_para.add_run('{{ findings }}')
findings_run.font.size = Pt(12)
findings_run.font.name = '仿宋_GB2312'

# 多个空行用于放置胸片图像
doc.add_paragraph()
doc.add_paragraph()

# 图片占位符段落 - 居中
img_para = doc.add_paragraph()
img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
img_placeholder_run = img_para.add_run('【胸片图像位置：{{ chest_xray.image_path }}】')
img_placeholder_run.font.size = Pt(10)
img_placeholder_run.font.name = '仿宋_GB2312'
img_placeholder_run.font.color.rgb = RGBColor(128, 128, 128)

# 空行
doc.add_paragraph()
doc.add_paragraph()

# 诊断意见标题
diagnosis_title_para = doc.add_paragraph()
diagnosis_title_run = diagnosis_title_para.add_run('诊断意见：')
diagnosis_title_run.font.size = Pt(12)
diagnosis_title_run.font.name = '仿宋_GB2312'

# 诊断意见内容
diagnosis_para = doc.add_paragraph()
diagnosis_run = diagnosis_para.add_run('{{ diagnosis }}')
diagnosis_run.font.size = Pt(12)
diagnosis_run.font.name = '仿宋_GB2312'

# 多个空行
for _ in range(3):
    doc.add_paragraph()

# 报告日期和医生
report_info_para = doc.add_paragraph()
report_info_run = report_info_para.add_run(
    '报告日期：{{ report_date }}      '
    '医生：{{ doctor_name }}      '
)
report_info_run.font.size = Pt(12)
report_info_run.font.name = '仿宋_GB2312'

# 注释
note_para = doc.add_paragraph()
note_run = note_para.add_run('注：本报告仅供临床医师参考。')
note_run.font.size = Pt(11)
note_run.font.name = '仿宋_GB2312'

# Save the template
doc.save('template.docx')
print("Medical report template created successfully!")
print(f"Total paragraphs: {len(doc.paragraphs)}")
